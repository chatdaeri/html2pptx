"""html2pptx 2단계 — 표 · 차트를 PowerPoint 네이티브로 만들고, 한글 글꼴(ea)을 지정하고, 글꼴을 임베딩한다.

    python3 native.py base.pptx natives.json out.pptx [--fontmap fontmap.json] [--no-embed]

base.pptx  : html2pptx.cjs 가 만든 글자 · 도형 장표 (표 · 차트 자리는 비어 있음)
natives    : 장표별 [{kind: table|chart, z, box, …}] — z = 그 앞에 그려진 도형 수 (같은 겹침 순서로 끼워 넣음)
단위: HTML px → EMU (1px = 9525 EMU = 0.75pt). 선 굵기 · 크기 보정 없음.
"""
import json
import os
import re
import sys
import tempfile
import zipfile

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from fontembed import embed  # noqa: E402

from lxml import etree  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.chart.data import CategoryChartData  # noqa: E402
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.opc.constants import RELATIONSHIP_TYPE as RT  # noqa: E402
from pptx.opc.package import Part  # noqa: E402
from pptx.enum.chart import XL_CHART_TYPE, XL_MARKER_STYLE, XL_TICK_MARK  # noqa: E402
from pptx.enum.dml import MSO_LINE_DASH_STYLE  # noqa: E402
from pptx.enum.shapes import MSO_CONNECTOR  # noqa: E402
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN  # noqa: E402
from pptx.oxml.ns import qn  # noqa: E402
from pptx.util import Emu, Pt  # noqa: E402

PX = 9525                     # EMU / px
NS = 'xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
FM = {'faces': {}, 'ea': {}, 'files': {}}


def xml(s):
    return etree.fromstring(f'<root {NS}>{s}</root>')[0]


def face(family, weight):
    """CSS family + weight → (PowerPoint 글꼴 이름, bold)."""
    t = FM['faces'].get(family)
    if not t:
        return family, weight >= 600
    best = min(t, key=lambda k: abs(int(k) - weight))
    f, b = t[best]
    return f, bool(b)


def defrpr(size_pt, color, fam='Inter', weight=500, tag='a:defRPr'):
    f, b = face(fam, weight)
    return (f'<{tag} sz="{round(size_pt * 100)}" b="{1 if b else 0}"><a:solidFill><a:srgbClr val="{color}"/></a:solidFill>'
            f'<a:latin typeface="{f}"/></{tag}>')


def txpr(size_pt, color, fam='Inter', weight=500):
    return f'<c:txPr><a:bodyPr/><a:lstStyle/><a:p><a:pPr>{defrpr(size_pt, color, fam, weight)}</a:pPr><a:endParaRPr lang="ko-KR"/></a:p></c:txPr>'


NOFILL = '<c:spPr><a:noFill/><a:ln><a:noFill/></a:ln></c:spPr>'

# ───────────────────────── 표 ─────────────────────────


def cluster(vals, tol=2.0):
    out = []
    for v in sorted(vals):
        if not out or v - out[-1] > tol:
            out.append(v)
    return out


def ln(tag, b):
    if not b:
        return f'<a:{tag} w="0"><a:noFill/></a:{tag}>'
    return (f'<a:{tag} w="{round(b[0] * 12700)}" cap="flat" cmpd="sng" algn="ctr"><a:solidFill><a:srgbClr val="{b[1]}"/></a:solidFill>'
            f'<a:prstDash val="solid"/><a:round/></a:{tag}>')


def set_run(run, o):
    f = run.font
    if o.get('fontSize'):
        f.size = Pt(o['fontSize'])
    if o.get('fontFace'):
        f.name = o['fontFace']
    f.bold = bool(o.get('bold'))
    if o.get('italic'):
        f.italic = True
    if o.get('color'):
        f.color.rgb = RGBColor.from_string(o['color'])
    if o.get('charSpacing'):
        run._r.get_or_add_rPr().set('spc', str(round(o['charSpacing'] * 100)))


def build_table(slide, spec):
    cells = spec['cells']
    xs = cluster([c['x'] for c in cells]) + [max(c['x'] + c['w'] for c in cells)]
    ys = cluster([c['y'] for c in cells]) + [max(c['y'] + c['h'] for c in cells)]
    near = lambda arr, v: min(range(len(arr)), key=lambda i: abs(arr[i] - v))  # noqa: E731
    nr, nc = len(ys) - 1, len(xs) - 1
    grid = {}
    for c in cells:
        r0, c0 = near(ys, c['y']), near(xs, c['x'])
        r1, c1 = max(near(ys, c['y'] + c['h']) - 1, r0), max(near(xs, c['x'] + c['w']) - 1, c0)
        grid[(r0, c0)] = (c, r1, c1)
    b = spec['box']
    gf = slide.shapes.add_table(nr, nc, Emu(int(b['x'] * 914400)), Emu(int(b['y'] * 914400)),
                                Emu(round((xs[-1] - xs[0]) * PX)), Emu(round((ys[-1] - ys[0]) * PX)))
    tbl = gf.table
    tblPr = tbl._tbl.tblPr
    for k in ('firstRow', 'bandRow'):
        tblPr.set(k, '0')
    for e in tblPr.findall(qn('a:tableStyleId')):
        tblPr.remove(e)
    for i in range(nc):
        tbl.columns[i].width = Emu(round((xs[i + 1] - xs[i]) * PX))
    for j in range(nr):
        tbl.rows[j].height = Emu(round((ys[j + 1] - ys[j]) * PX))
    # 칸 경계선: 위 = 이 칸의 윗선(합계 · 강조 열) 우선, 없으면 위 칸 아랫선 / 첫 행은 표 윗선
    owner = {}
    for (r0, c0), (c, r1, c1) in grid.items():
        for r in range(r0, r1 + 1):
            for cc in range(c0, c1 + 1):
                owner[(r, cc)] = c
    for (r0, c0), (c, r1, c1) in grid.items():
        cell = tbl.cell(r0, c0)
        if r1 > r0 or c1 > c0:
            cell.merge(tbl.cell(r1, c1))
        above = owner.get((r0 - 1, c0))
        top = c['border']['t'] or (above['border']['b'] if above else spec.get('top'))
        if r0 == 0 and not c['border']['t']:
            top = spec.get('top')
        bottom = c['border']['b']
        below = owner.get((r1 + 1, c0))
        if below and below['border']['t']:
            bottom = below['border']['t']
        cell.margin_left, cell.margin_right = Emu(round(c['pad'][0] * 12700)), Emu(round(c['pad'][1] * 12700))
        cell.margin_top, cell.margin_bottom = Emu(round(c['pad'][2] * 12700)), Emu(round(c['pad'][3] * 12700))
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        if c['fill']:
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor.from_string(c['fill'])
        else:
            cell.fill.background()
        tcPr = cell._tc.get_or_add_tcPr()
        for e in [x for x in tcPr if etree.QName(x).localname in ('lnL', 'lnR', 'lnT', 'lnB')]:
            tcPr.remove(e)
        for i, tag in enumerate(('lnL', 'lnR', 'lnT', 'lnB')):
            tcPr.insert(i, xml(ln(tag, {'lnL': c['border']['l'], 'lnR': c['border']['r'], 'lnT': top, 'lnB': bottom}[tag])))
        tf = cell.text_frame
        tf.word_wrap = True
        for k, para in enumerate(c['paras'] or [{'style': {}, 'runs': []}]):
            p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
            st = para['style']
            p.alignment = {'right': PP_ALIGN.RIGHT, 'center': PP_ALIGN.CENTER}.get(st.get('align'), PP_ALIGN.LEFT)
            if st.get('lineSpacing'):
                p.line_spacing = Pt(st['lineSpacing'])
            for ro in para['runs']:
                set_run(p.add_run(), {**st, **ro['options'], 'text': ro['text']})
                p.runs[-1].text = ro['text']
    return gf

# ───────────────────────── 차트 ─────────────────────────


def chart_common(ch, spec, inner):
    """글꼴 기본값 · 배경 없음 · 범례/제목 없음 · 플롯 영역 수동 배치(HTML 그림 영역과 같은 자리)."""
    pal = spec['palette']
    ch.has_legend = False
    ch.has_title = False
    cs = ch._chartSpace
    cs.find(qn('c:chart')).addnext(xml(NOFILL))
    rc = cs.find(qn('c:roundedCorners'))
    if rc is None:   # 스키마 순서: date1904 → lang → roundedCorners
        prev = [e for e in cs if etree.QName(e).localname in ('date1904', 'lang')]
        rc = xml('<c:roundedCorners val="0"/>')
        (prev[-1].addnext(rc) if prev else cs.insert(0, rc))
    else:
        rc.set('val', '0')
    for old in cs.findall(qn('c:txPr')):   # python-pptx 기본 txPr 교체 (중복이면 PowerPoint 가 '복구' 요구)
        cs.remove(old)
    cs.find(qn('c:spPr')).addnext(xml(txpr(9, pal['charcoal'])))
    pa = cs.find(qn('c:chart')).find(qn('c:plotArea'))
    lay = pa.find(qn('c:layout'))
    if lay is None:
        lay = xml('<c:layout/>')
        pa.insert(0, lay)
    for e in list(lay):
        lay.remove(e)
    W, H = spec['size']['w'], spec['size']['h']
    x, y, w, h = inner
    lay.append(xml(f'<c:manualLayout><c:layoutTarget val="inner"/><c:xMode val="edge"/><c:yMode val="edge"/>'
                   f'<c:x val="{x / W:.5f}"/><c:y val="{y / H:.5f}"/><c:w val="{w / W:.5f}"/><c:h val="{h / H:.5f}"/></c:manualLayout>'))
    ext = pa.find(qn('c:extLst'))
    sp = xml(NOFILL)
    (ext.addprevious(sp) if ext is not None else pa.append(sp))
    return pa


def frame_fit(gf, spec, inner):
    """그림 영역이 차트 틀 밖으로 나가면 틀을 넓힌다 (PowerPoint 는 틀 밖 플롯 영역을 안으로 줄여 버림)."""
    x, y, w, h = inner
    W, H = spec['size']['w'], spec['size']['h']
    l, t, r, btm = min(0, x), min(0, y), max(0, x + w - W), max(0, y + h - H)
    if l or t or r or btm:
        gf.left, gf.top = Emu(gf.left - round(-l * PX)), Emu(gf.top - round(-t * PX))
        gf.width, gf.height = Emu(gf.width + round((r - l) * PX)), Emu(gf.height + round((btm - t) * PX))
        spec = {**spec, 'size': {'w': W - l + r, 'h': H - t + btm}}
        inner = (x - l, y - t, w, h)
    return spec, inner


def ser_el(ch, i):
    return ch._chartSpace.findall('.//' + qn('c:ser'))[i]


def put_dlbls(ser, dlbls_xml):
    for e in ser.findall(qn('c:dLbls')):
        ser.remove(e)
    ser.find(qn('c:cat')).addprevious(xml(dlbls_xml))


SHOW = lambda val=0, ser=0, sep=None: (  # noqa: E731
    f'<c:showLegendKey val="0"/><c:showVal val="{val}"/><c:showCatName val="0"/><c:showSerName val="{ser}"/>'
    f'<c:showPercent val="0"/><c:showBubbleSize val="0"/>' + (f'<c:separator>{sep}</c:separator>' if sep else ''))


def dlbl(idx, size, color, fam, weight, pos, val=1, ser=0, sep=None):
    return (f'<c:dLbl><c:idx val="{idx}"/>{NOFILL}{txpr(size, color, fam, weight)}<c:dLblPos val="{pos}"/>'
            f'{SHOW(val, ser, sep)}</c:dLbl>')


def move_to_line(ch, ser_index, spPr, labels_xml=None, marker='<c:marker><c:symbol val="none"/></c:marker>', ax_ids=None):
    """막대 차트의 계열 하나를 꺾은선으로 옮긴다 (평균선 · 합계 레이블 · 이중축 선) — 데이터 편집에도 그대로 남음.
    ax_ids 가 없으면 막대와 같은 축, 있으면 그 축(보조 축)을 쓴다."""
    pa = ch._chartSpace.find('.//' + qn('c:plotArea'))
    bc = pa.find(qn('c:barChart'))
    s = bc.findall(qn('c:ser'))[ser_index]
    bc.remove(s)
    for tag in ('c:invertIfNegative', 'c:shape', 'c:dPt', 'c:dLbls'):
        for e in s.findall(qn(tag)):
            s.remove(e)
    for e in s.findall(qn('c:spPr')):
        s.remove(e)
    s.find(qn('c:tx')).addnext(xml(spPr))
    s.find(qn('c:spPr')).addnext(xml(marker))
    if labels_xml:
        s.find(qn('c:cat')).addprevious(xml(labels_xml))
    s.append(xml('<c:smooth val="0"/>'))
    lc = xml('<c:lineChart><c:grouping val="standard"/><c:varyColors val="0"/></c:lineChart>')
    lc.append(s)
    lc.append(xml('<c:marker val="1"/>'))
    for ax in ax_ids or [a.get('val') for a in bc.findall(qn('c:axId'))]:
        lc.append(xml(f'<c:axId val="{ax}"/>'))
    bc.addnext(lc)
    return lc


def add_secondary_axes(ch, cat_id, val_id, lo, hi, step, fmt, pal):
    """오른쪽 보조 값 축 + 숨긴 보조 항목 축 (이중축). 막대 축 바로 뒤에 둔다 (스키마 순서: 차트들 → 축들)."""
    pa = ch._chartSpace.find('.//' + qn('c:plotArea'))
    last_ax = [e for e in pa if etree.QName(e).localname in ('catAx', 'valAx', 'dateAx', 'serAx')][-1]
    cat = xml(f'<c:catAx><c:axId val="{cat_id}"/><c:scaling><c:orientation val="minMax"/></c:scaling><c:delete val="1"/>'
              f'<c:axPos val="b"/><c:majorTickMark val="none"/><c:minorTickMark val="none"/><c:tickLblPos val="nextTo"/>'
              f'<c:crossAx val="{val_id}"/><c:crosses val="autoZero"/><c:auto val="1"/><c:lblAlgn val="ctr"/><c:lblOffset val="100"/>'
              f'<c:noMultiLvlLbl val="0"/></c:catAx>')
    val = xml(f'<c:valAx><c:axId val="{val_id}"/><c:scaling><c:orientation val="minMax"/><c:max val="{hi}"/><c:min val="{lo}"/></c:scaling>'
              f'<c:delete val="0"/><c:axPos val="r"/><c:numFmt formatCode="{fmt}" sourceLinked="0"/>'
              f'<c:majorTickMark val="none"/><c:minorTickMark val="none"/><c:tickLblPos val="nextTo"/>'
              f'<c:spPr><a:ln><a:noFill/></a:ln></c:spPr>{txpr(9, pal["stone"])}'
              f'<c:crossAx val="{cat_id}"/><c:crosses val="max"/><c:crossBetween val="between"/><c:majorUnit val="{step}"/></c:valAx>')
    last_ax.addnext(cat)
    cat.addnext(val)


def axis_style(ax, pal, show_line=True, labels=True, size=9, color=None):
    ax.major_tick_mark = XL_TICK_MARK.NONE
    ax.minor_tick_mark = XL_TICK_MARK.NONE
    if show_line:
        ax.format.line.color.rgb = RGBColor.from_string(pal['ink'])
        ax.format.line.width = Pt(0.75)
    else:
        ax.format.line.fill.background()
    if labels:
        tl = ax.tick_labels.font
        tl.size = Pt(size)
        f, b = face('Inter', 500)
        tl.name, tl.bold = f, b
        tl.color.rgb = RGBColor.from_string(color or pal['charcoal'])


def xl_fmt(f, default='#,##0'):
    """charts.js 표시 형식 → 엑셀 형식. 끝의 % 는 글자로 붙인다 (charts.js 처럼 값 그대로 · 엑셀 % 처럼 ×100 하지 않음)."""
    f = f or default
    return re.sub(r'(?<!")%$', '"%"', f)


def fmt_code(c, default='#,##0'):
    return xl_fmt(c.get('format'), default).replace('"', '&quot;')


def build_chart(slide, spec):
    c, pal, g, pl = spec['chart'], spec['palette'], spec['geom'], spec['plot']
    b = spec['box']
    X, Y = Emu(int(b['x'] * 914400)), Emu(int(b['y'] * 914400))
    CX, CY = Emu(round(spec['size']['w'] * PX)), Emu(round(spec['size']['h'] * PX))
    t = c['type']
    cd = CategoryChartData(number_format=xl_fmt(c.get('format')))
    extra = []   # 차트 위에 얹는 도형 (사건 세로선)

    if t in ('bar', 'stack'):
        cd.categories = c['categories']
        series = c['series'] if t == 'stack' else [{'name': c.get('name', '값'), 'values': c['values']}]
        for se in series:
            cd.add_series(se['name'], se['values'])
        if t == 'stack':
            cd.add_series('합계', [sum(se['values'][i] for se in series) for i in range(len(c['categories']))])
        if c.get('ref'):
            cd.add_series(c['ref'].get('name', '기준'), [c['ref']['value']] * len(c['categories']))
        gf = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_STACKED if t == 'stack' else XL_CHART_TYPE.COLUMN_CLUSTERED, X, Y, CX, CY, cd)
        ch = gf.chart
        n = len(c['categories'])
        cw = pl['w'] / n
        plot = ch.plots[0]
        plot.gap_width = max(0, min(500, round((cw - g['barW']) / g['barW'] * 100)))
        plot.overlap = 100 if t == 'stack' else 0
        plot.vary_by_categories = False
        ramp = [pal['charcoal'], pal['ash'], pal['pebble']]
        for k, se in enumerate(series):
            s = ch.series[k]
            s.format.fill.solid()
            s.format.fill.fore_color.rgb = RGBColor.from_string(ramp[k % 3] if t == 'stack' else pal['ash'])
            if t == 'stack':   # 구간 사이 2px 틈 = 흰 테두리
                s.format.line.color.rgb = RGBColor.from_string(pal['paper'])
                s.format.line.width = Pt(1.5)
        key = c.get('key')
        if t == 'bar' and key is not None:
            pt = ch.series[0].points[key]
            pt.format.fill.solid()
            pt.format.fill.fore_color.rgb = RGBColor.from_string(pal['accent'])
        if t == 'stack' and isinstance(key, list):
            pt = ch.series[key[0]].points[key[1]]
            pt.format.fill.solid()
            pt.format.fill.fore_color.rgb = RGBColor.from_string(pal['accent'])
            pt.format.line.color.rgb = RGBColor.from_string(pal['paper'])
            pt.format.line.width = Pt(1.5)
        va = ch.value_axis
        va.minimum_scale, va.maximum_scale = 0, c['max']
        va.has_major_gridlines = False
        va.visible = False
        axis_style(ch.category_axis, pal)
        num = f'<c:numFmt formatCode="{fmt_code(c)}" sourceLinked="0"/>'
        if t == 'bar':
            keylbl = dlbl(key, 15, pal['ink'], 'General Sans', 700, 'outEnd') if key is not None else ''
            put_dlbls(ser_el(ch, 0), f'<c:dLbls>{keylbl}{num}{NOFILL}{txpr(10.5, pal["ink"], "Inter", 600)}<c:dLblPos val="outEnd"/>{SHOW(1)}</c:dLbls>')
        idx = len(series)
        if t == 'stack':   # 합계 → 선 없는 꺾은선 + 위쪽 레이블
            move_to_line(ch, idx, '<c:spPr><a:ln><a:noFill/></a:ln></c:spPr>',
                         f'<c:dLbls>{num}{NOFILL}{txpr(10.5, pal["ink"], "Inter", 600)}<c:dLblPos val="t"/>{SHOW(1)}</c:dLbls>')
        if c.get('ref'):   # 평균 · 목표 → 점선 꺾은선 (같은 데이터)
            move_to_line(ch, len(ch._chartSpace.find('.//' + qn('c:barChart')).findall(qn('c:ser'))) - 1,
                         f'<c:spPr><a:ln w="9525"><a:solidFill><a:srgbClr val="{pal["charcoal"]}"/></a:solidFill><a:prstDash val="dash"/></a:ln></c:spPr>')
        chart_common(ch, *frame_fit(gf, spec, (pl['x'], pl['y'], pl['w'], pl['h'])))

    elif t == 'hbar':
        cd.categories = c['categories']
        cd.add_series(c.get('name', '값'), c['values'])
        gf = slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, X, Y, CX, CY, cd)
        ch = gf.chart
        plot = ch.plots[0]
        plot.gap_width = max(0, min(500, round((g['pitch'] - g['barH']) / g['barH'] * 100)))
        plot.vary_by_categories = False
        s = ch.series[0]
        s.format.fill.solid()
        s.format.fill.fore_color.rgb = RGBColor.from_string(pal['ash'])
        if c.get('key') is not None:
            pt = s.points[c['key']]
            pt.format.fill.solid()
            pt.format.fill.fore_color.rgb = RGBColor.from_string(pal['accent'])
        ca = ch.category_axis
        ca.reverse_order = True       # 첫 항목이 맨 위
        axis_style(ca, pal, show_line=False, size=10.5)
        va = ch.value_axis
        va.minimum_scale, va.maximum_scale = 0, c['max']
        va.has_major_gridlines = False
        va.visible = False
        put_dlbls(ser_el(ch, 0), f'<c:dLbls><c:numFmt formatCode="{fmt_code(c)}" sourceLinked="0"/>{NOFILL}'
                                 f'{txpr(10.5, pal["ink"], "Inter", 600)}<c:dLblPos val="outEnd"/>{SHOW(1)}</c:dLbls>')
        n = len(c['categories'])
        top = g['c0'] - g['pitch'] / 2
        chart_common(ch, *frame_fit(gf, spec, (g['x0'], top, g['fullW'], g['pitch'] * n)))

    elif t == 'line':
        cd = CategoryChartData(number_format=xl_fmt(c.get('format'), '0'))
        cd.categories = c['categories']
        for se in c['series']:
            cd.add_series(se['name'], se['values'])
        gf = slide.shapes.add_chart(XL_CHART_TYPE.LINE, X, Y, CX, CY, cd)
        ch = gf.chart
        last = len(c['categories']) - 1
        for k, se in enumerate(c['series']):
            s = ch.series[k]
            s.smooth = False
            s.marker.style = XL_MARKER_STYLE.NONE
            ln_ = s.format.line
            style = se.get('style', 'base')
            ln_.color.rgb = RGBColor.from_string(pal['accent'] if style == 'key' else pal['ash'] if style == 'dash' else pal['stone'])
            ln_.width = Pt(2.25 if style == 'key' else 1.125)
            if style == 'dash':
                ln_.dash_style = MSO_LINE_DASH_STYLE.DASH
            if style == 'key':
                m = s.points[last].marker
                m.style, m.size = XL_MARKER_STYLE.CIRCLE, 9
                m.format.fill.solid()
                m.format.fill.fore_color.rgb = RGBColor.from_string(pal['paper'])
                m.format.line.color.rgb = RGBColor.from_string(pal['accent'])
                m.format.line.width = Pt(1.875)
            key = style == 'key'
            put_dlbls(ser_el(ch, k), '<c:dLbls>' + dlbl(last, 9, pal['ink'] if key else pal['charcoal'], 'Inter', 600 if key else 500, 'r', 1, 1, ' ')
                      + f'{NOFILL}{SHOW(0)}</c:dLbls>')
        va = ch.value_axis
        va.minimum_scale, va.maximum_scale, va.major_unit = c['min'], c['max'], c['step']
        va.has_major_gridlines = True
        va.major_gridlines.format.line.color.rgb = RGBColor.from_string(pal['pebble'])
        va.major_gridlines.format.line.width = Pt(0.75)
        axis_style(va, pal, show_line=False, color=pal['stone'])
        va.tick_labels.number_format, va.tick_labels.number_format_is_linked = xl_fmt(c.get('format'), '0'), False
        ca = ch.category_axis
        axis_style(ca, pal, color=pal['stone'])
        cax = ch._chartSpace.find('.//' + qn('c:catAx'))
        if c.get('labelEvery', 1) > 1:
            anchor = cax.find(qn('c:noMultiLvlLbl'))
            el = xml(f'<c:tickLblSkip val="{c["labelEvery"]}"/>')
            (anchor.addprevious(el) if anchor is not None else cax.append(el))
        cb = ch._chartSpace.find('.//' + qn('c:valAx')).find(qn('c:crossBetween'))
        if cb is not None:
            cb.set('val', 'midCat')   # 첫 · 끝 점이 플롯 가장자리 (HTML 과 같게)
        chart_common(ch, *frame_fit(gf, spec, (pl['x'], pl['y'], pl['w'], pl['h'])))
        if c.get('note'):
            i = c['note']['index']
            k = next((se for se in c['series'] if se.get('style') == 'key'), c['series'][0])
            px_ = pl['x'] + pl['w'] * i / max(last, 1)
            py_ = pl['y'] + pl['h'] * (1 - (k['values'][i] - c['min']) / (c['max'] - c['min']))
            extra.append(('vline', px_, py_, pl['y'] + pl['h']))

    elif t == 'dual':   # 막대(왼쪽 축) + 꺾은선(오른쪽 보조 축) — 엑셀 데이터 한 표에 두 계열
        bs, ls = c['bar'], c['line']
        cd = CategoryChartData(number_format=xl_fmt(bs.get('format')))
        cd.categories = c['categories']
        cd.add_series(bs['name'], bs['values'])
        cd.add_series(ls['name'], ls['values'], number_format=xl_fmt(ls.get('format'), '0'))
        gf = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, X, Y, CX, CY, cd)
        ch = gf.chart
        cw = pl['w'] / len(c['categories'])
        plot = ch.plots[0]
        plot.gap_width = max(0, min(500, round((cw - g['barW']) / g['barW'] * 100)))
        plot.overlap = 0
        plot.vary_by_categories = False
        s = ch.series[0]
        s.format.fill.solid()
        s.format.fill.fore_color.rgb = RGBColor.from_string(pal['ash'])
        key = c.get('key')
        if key is not None:
            pt = s.points[key]
            pt.format.fill.solid()
            pt.format.fill.fore_color.rgb = RGBColor.from_string(pal['accent'])
        va = ch.value_axis
        va.minimum_scale, va.maximum_scale, va.major_unit = 0, bs['max'], bs['step']
        va.has_major_gridlines = True
        va.major_gridlines.format.line.color.rgb = RGBColor.from_string(pal['pebble'])
        va.major_gridlines.format.line.width = Pt(0.75)
        axis_style(va, pal, show_line=False, color=pal['stone'])
        va.tick_labels.number_format, va.tick_labels.number_format_is_linked = xl_fmt(bs.get('format')), False
        axis_style(ch.category_axis, pal)
        if c.get('labelEvery', 1) > 1:
            cax = ch._chartSpace.find('.//' + qn('c:catAx'))
            anchor = cax.find(qn('c:noMultiLvlLbl'))
            el = xml(f'<c:tickLblSkip val="{c["labelEvery"]}"/>')
            (anchor.addprevious(el) if anchor is not None else cax.append(el))
        if bs.get('labels'):
            keylbl = dlbl(key, 15, pal['ink'], 'General Sans', 700, 'outEnd') if key is not None else ''
            put_dlbls(ser_el(ch, 0), f'<c:dLbls>{keylbl}<c:numFmt formatCode="{fmt_code(bs)}" sourceLinked="0"/>{NOFILL}'
                                     f'{txpr(10.5, pal["ink"], "Inter", 600)}<c:dLblPos val="outEnd"/>{SHOW(1)}</c:dLbls>')
        ink = f'<a:solidFill><a:srgbClr val="{pal["ink"]}"/></a:solidFill>'
        lbl = (f'<c:dLbls><c:numFmt formatCode="{fmt_code(ls, "0")}" sourceLinked="0"/>{NOFILL}{txpr(9, pal["ink"], "Inter", 600)}'
               f'<c:dLblPos val="t"/>{SHOW(1)}</c:dLbls>') if ls.get('labels', True) else None
        move_to_line(ch, 1, f'<c:spPr><a:ln w="19050" cap="rnd">{ink}<a:round/></a:ln></c:spPr>', lbl,
                     marker=f'<c:marker><c:symbol val="circle"/><c:size val="6"/><c:spPr><a:solidFill><a:srgbClr val="{pal["paper"]}"/></a:solidFill>'
                            f'<a:ln w="19050">{ink}</a:ln></c:spPr></c:marker>',
                     ax_ids=['70010001', '70010002'])
        add_secondary_axes(ch, '70010001', '70010002', ls['min'], ls['max'], ls['step'], fmt_code(ls, '0'), pal)
        chart_common(ch, *frame_fit(gf, spec, (pl['x'], pl['y'], pl['w'], pl['h'])))

    elif t in ('donut', 'ring'):
        ring = t == 'ring'
        cd = CategoryChartData(number_format='0')
        cd.categories = ['값', '나머지'] if ring else c['categories']
        cd.add_series(c.get('name', '비중'), [c['value'], 100 - c['value']] if ring else c['values'])
        gf = slide.shapes.add_chart(XL_CHART_TYPE.DOUGHNUT, X, Y, CX, CY, cd)
        ch = gf.chart
        plot = ch.plots[0]
        plot.vary_by_categories = True
        ramp = [pal['ink'], pal['stone'], pal['ash'], pal['pebble'], pal['charcoal']]
        r = 0
        s = ch.series[0]
        for i in range(2 if ring else len(c['values'])):
            if ring:
                col = (pal['accent'] if c.get('key') else pal['ink']) if i == 0 else pal['pebble']
            elif c.get('key') == i:
                col = pal['accent']
            else:
                col = ramp[r % len(ramp)]
                r += 1
            pt = s.points[i]
            pt.format.fill.solid()
            pt.format.fill.fore_color.rgb = RGBColor.from_string(col)
            if ring:
                pt.format.line.fill.background()
            else:   # 조각 사이 틈
                pt.format.line.color.rgb = RGBColor.from_string(pal['paper'])
                pt.format.line.width = Pt(1.5)
        dc = ch._chartSpace.find('.//' + qn('c:doughnutChart'))
        hs, fa = dc.find(qn('c:holeSize')), dc.find(qn('c:firstSliceAng'))
        stroke = 4 if ring else 6
        outer, inner = 15.915 + stroke / 2, 15.915 - stroke / 2
        if hs is not None:
            hs.set('val', str(round(inner / outer * 100)))
        if fa is not None:
            fa.set('val', '0')
        W = spec['size']['w']
        d = W * outer / 21           # 링 바깥 지름 (viewBox 42 기준)
        chart_common(ch, *frame_fit(gf, spec, ((W - d) / 2, (spec['size']['h'] - d) / 2, d, d)))
    else:
        raise ValueError(f'알 수 없는 차트 종류: {t}')
    return gf, extra


def add_extra(slide, spec, extra):
    b = spec['box']
    pal = spec['palette']
    shapes = []
    for kind, x, y1, y2 in extra:
        if kind == 'vline':
            X = Emu(int(b['x'] * 914400) + round(x * PX))
            ln_ = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, X, Emu(int(b['y'] * 914400) + round(y1 * PX)),
                                             X, Emu(int(b['y'] * 914400) + round(y2 * PX)))
            ln_.line.color.rgb = RGBColor.from_string(pal['charcoal'])
            ln_.line.width = Pt(0.75)
            ln_.line.dash_style = MSO_LINE_DASH_STYLE.DASH
            shapes.append(ln_)
    return shapes

# ───────────────────────── 그림 ─────────────────────────

SVG_EXT = '{96DAC541-7B7A-43D3-8B79-37D633B846F1}'   # PowerPoint 2016+ SVG 확장 (대체 PNG + SVG 원본)


def attach_svg(prs, slide, name, svg_text):
    """대체 PNG 로 들어간 그림에 SVG 원본을 붙인다 → PowerPoint 에서 선명한 벡터 그림 · '도형으로 변환' 가능."""
    pic = next((sh for sh in slide.shapes if sh.name == name), None)
    if pic is None:
        raise ValueError(f'SVG 그림 {name} 을 찾지 못했습니다')
    pkg = prs.part.package
    part = Part(pkg.next_partname('/ppt/media/vector%d.svg'), 'image/svg+xml', pkg, svg_text.encode('utf8'))
    rid = slide.part.relate_to(part, RT.IMAGE)
    blip = pic._element.find('.//' + qn('a:blip'))
    for e in blip.findall(qn('a:extLst')):
        blip.remove(e)
    blip.append(etree.fromstring(
        f'<a:extLst xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><a:ext uri="{SVG_EXT}">'
        f'<asvg:svgBlip xmlns:asvg="http://schemas.microsoft.com/office/drawing/2016/SVG/main" '
        f'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" r:embed="{rid}"/></a:ext></a:extLst>'))


def round_pic(slide, name, r, w, h):
    """그림 모서리를 CSS border-radius 와 같게 (PowerPoint 둥근 사각형 그림 틀)."""
    pic = next((sh for sh in slide.shapes if sh.name == name), None)
    if pic is None:
        return
    geom = pic._element.find('.//' + qn('a:prstGeom'))
    geom.set('prst', 'roundRect')
    for e in list(geom):
        geom.remove(e)
    geom.append(xml(f'<a:avLst><a:gd name="adj" fmla="val {min(50000, round(r / min(w, h) * 100000))}"/></a:avLst>'))

# ───────────────────────── 글꼴 ─────────────────────────


def fix_ea(x):
    """라틴 글꼴마다 한글(동아시아) 글꼴을 짝지어 지정 — PowerPoint 는 한글 글자에 ea 글꼴을 쓴다."""
    def rep(m):
        ea = FM['ea'].get(m.group(2))
        return m.group(1) + (f'<a:ea typeface="{ea}"/>' if ea else (m.group(3) or ''))
    return re.sub(r'(<a:latin typeface="([^"]+)"[^>]*/>)(<a:ea [^>]*/>)?', rep, x)


def finish(path_in, path_out, fm_dir, do_embed):
    zin = zipfile.ZipFile(path_in)
    tmp = path_out + '.tmp'
    used = set()
    with zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as z:
        for n in zin.namelist():
            d = zin.read(n)
            if re.match(r'ppt/(slides|charts|slideLayouts|slideMasters)/.*\.xml$', n):
                s = fix_ea(d.decode('utf8'))
                used |= set(re.findall(r'typeface="([^"+][^"]*)"', s))
                d = s.encode('utf8')
            z.writestr(n, d)
    if do_embed and FM['files']:
        faces = {}
        for name in sorted(used & set(FM['files'])):
            faces[name] = {st: open(os.path.join(fm_dir, f), 'rb').read() for st, f in FM['files'][name].items()}
        embed(tmp, path_out, faces)
        os.remove(tmp)
        print(f'   글꼴 임베딩 {len(faces)}종: ' + ', '.join(faces))
    else:
        os.replace(tmp, path_out)


# ───────────────────────── 영상 자동 재생 ─────────────────────────


def renumber_ids(slide):
    """장표 안 도형 번호(cNvPr id)를 1부터 겹치지 않게 다시 매긴다.
    PptxGenJS 는 영상(addMedia)에 다른 도형과 같은 번호를 줄 때가 있고, 영상 자동 재생 타이밍이 그 번호를 가리키면
    PowerPoint 가 '복구' 창을 띄운다 (2026-10-17 실제 발생)."""
    A = '{http://schemas.openxmlformats.org/presentationml/2006/main}'
    for k, el in enumerate(slide.shapes._spTree.iter(A + 'cNvPr'), start=1):
        el.set('id', str(k))


def video_timing(slide, videos):
    """슬라이드쇼에서 장표가 열리면 영상이 바로 재생되도록 <p:timing> 을 넣는다 (PowerPoint '시작: 자동' · '반복 재생' 과 같은 구조).
    소리 없는 모션그래픽이므로 음량 0. 장표에 이미 timing 이 있으면 건드리지 않는다."""
    sld = slide._element
    P = '{http://schemas.openxmlformats.org/presentationml/2006/main}'
    if sld.find(P + 'timing') is not None:
        print(f'   경고: 장표에 이미 애니메이션이 있어 영상 자동 재생을 넣지 않았습니다')
        return 0
    ids = {sh.name: sh.shape_id for sh in slide.shapes}
    auto = [(ids[v['name']], v) for v in videos if v['name'] in ids and v.get('autoplay', True)]
    if not auto:
        return 0
    n = [6]   # cTn id 1~5 는 고정 틀
    def nid():
        n[0] += 1
        return n[0]
    calls = ''.join(
        f'<p:par><p:cTn id="{nid()}" presetID="1" presetClass="mediacall" presetSubtype="0" fill="hold" '
        f'nodeType="{"afterEffect" if k == 0 else "withEffect"}"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
        f'<p:cmd type="call" cmd="playFrom(0.0)"><p:cBhvr><p:cTn id="{nid()}" dur="1" fill="hold"/>'
        f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cBhvr></p:cmd></p:childTnLst></p:cTn></p:par>'
        for k, (spid, _) in enumerate(auto))
    media = ''.join(
        f'<p:video><p:cMediaNode vol="0" mute="1"><p:cTn id="{nid()}"{" repeatCount=\"indefinite\"" if v.get("loop", True) else ""} fill="hold" display="0">'
        f'<p:stCondLst><p:cond delay="indefinite"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cMediaNode></p:video>'
        for spid, v in auto)
    timing = xml(
        '<p:timing xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"><p:tnLst><p:par>'
        '<p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>'
        '<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>'
        '<p:par><p:cTn id="3" fill="hold"><p:stCondLst><p:cond delay="indefinite"/>'
        '<p:cond evt="onBegin" delay="0"><p:tn val="2"/></p:cond></p:stCondLst><p:childTnLst>'
        '<p:par><p:cTn id="4" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
        + calls +
        '</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn>'
        '<p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
        '<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq>'
        + media +
        '</p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>')
    ext = sld.find(P + 'extLst')   # 순서: cSld · clrMapOvr · transition · timing · extLst
    if ext is not None:
        ext.addprevious(timing)
    else:
        sld.append(timing)
    return len(auto)


def main():
    a = sys.argv[1:]
    base, nat, out = a[0], a[1], a[2]
    fm_dir = None
    if '--fontmap' in a:
        fp = a[a.index('--fontmap') + 1]
        FM.update(json.load(open(fp)))
        fm_dir = os.path.dirname(os.path.abspath(fp))
    prs = Presentation(base)
    natives = json.load(open(nat))
    count = {'table': 0, 'chart': 0, 'svg': 0, 'video': 0}
    for entry in natives:
        slide = prs.slides[entry['slide']]
        for v in entry.get('svgs', []):
            attach_svg(prs, slide, v['name'], v['svg'])
            count['svg'] += 1
        for pc in entry.get('pics', []):
            round_pic(slide, pc['name'], pc['r'], pc['w'], pc['h'])
        renumber_ids(slide)   # 표 · 차트를 넣은 뒤가 아니라 여기서 — 아래 표 · 차트는 python-pptx 가 겹치지 않는 번호를 줌
        if entry.get('videos'):
            count['video'] += video_timing(slide, entry['videos'])
        tree = slide.shapes._spTree
        offset = 0
        for spec in sorted(entry['natives'], key=lambda s: s['z']):
            if spec['kind'] == 'table':
                els = [build_table(slide, spec)]
            else:
                gf, extra = build_chart(slide, spec)
                els = [gf] + add_extra(slide, spec, extra)
            for e in els:   # HTML 과 같은 겹침 순서 자리로 이동
                el = e._element
                tree.remove(el)
                tree.insert(2 + spec['z'] + offset, el)
                offset += 1
            count[spec['kind']] += 1
    fd, mid = tempfile.mkstemp(suffix='.pptx')
    os.close(fd)
    prs.save(mid)
    finish(mid, out, fm_dir, '--no-embed' not in a)
    os.remove(mid)
    print(f'   표 {count["table"]}개 · 차트 {count["chart"]}개 네이티브 생성' + (f' · SVG 벡터 그림 {count["svg"]}개' if count['svg'] else '')
          + (f' · 영상 자동 재생 {count["video"]}개' if count['video'] else ''))


if __name__ == '__main__':
    main()
